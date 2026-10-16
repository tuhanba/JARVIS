"""Stage F: document reading with locations, limits and broken input (T26/T27 parts)."""
import threading

from jarvis_core.araclar import Baglam, KAYIT


def _b():
    return Baglam(threading.Event())


def _pdf(yol, sayfalar):
    """Minimal text PDF writer (no extra dependency) for fixtures."""
    nesneler, sayfa_ids = [], []
    nesneler.append(b'<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>')
    for metin in sayfalar:
        akis = f'BT /F1 12 Tf 72 720 Td ({metin}) Tj ET'.encode('latin-1')
        nesneler.append(b'<< /Length %d >>\nstream\n' % len(akis) + akis + b'\nendstream')
        nesneler.append(None)
        sayfa_ids.append(len(nesneler))
    kids_no = len(nesneler) + 1
    for i, no in enumerate(sayfa_ids):
        nesneler[no - 1] = (b'<< /Type /Page /Parent %d 0 R /MediaBox [0 0 612 792] /Contents %d 0 R '
                            b'/Resources << /Font << /F1 1 0 R >> >> >>' % (kids_no, no - 1))
    nesneler.append(b'<< /Type /Pages /Kids [%s] /Count %d >>' % (
        b' '.join(b'%d 0 R' % n for n in sayfa_ids), len(sayfa_ids)))
    nesneler.append(b'<< /Type /Catalog /Pages %d 0 R >>' % kids_no)
    cikti, ofsetler = b'%PDF-1.4\n', []
    for i, n in enumerate(nesneler, 1):
        ofsetler.append(len(cikti))
        cikti += b'%d 0 obj\n' % i + n + b'\nendobj\n'
    xref = len(cikti)
    cikti += b'xref\n0 %d\n0000000000 65535 f \n' % (len(nesneler) + 1)
    cikti += b''.join(b'%010d 00000 n \n' % o for o in ofsetler)
    cikti += b'trailer\n<< /Size %d /Root %d 0 R >>\nstartxref\n%d\n%%%%EOF\n' % (len(nesneler) + 1, len(nesneler), xref)
    yol.write_bytes(cikti)


def test_pdf_sayfa_konumlariyla_okunur(ev):
    _pdf(ev / 'Downloads' / 'fatura.pdf', ['Fatura no 1001', 'Toplam 250 TL'])
    s = KAYIT['belge_oku'].calistir(_b(), yol='İndirilenler/fatura.pdf')
    assert s.basarili and '[sayfa 1] Fatura no 1001' in s.cikti and '[sayfa 2] Toplam 250 TL' in s.cikti
    assert len(s.veri['sha256']) == 64


def test_docx_ve_xlsx(ev):
    import docx
    import openpyxl
    d = docx.Document()
    d.add_paragraph('Toplantı notları')
    d.add_paragraph('Karar: bütçe onaylandı.')
    d.save(ev / 'Documents' / 'not.docx')
    k = openpyxl.Workbook()
    k.active.title = 'Giderler'
    k.active.append(['Tarih', 'Tutar'])
    k.active.append(['2026-10-01', 120.5])
    k.save(ev / 'Documents' / 'gider.xlsx')
    s1 = KAYIT['belge_oku'].calistir(_b(), yol='Belgeler/not.docx')
    s2 = KAYIT['belge_oku'].calistir(_b(), yol='Belgeler/gider.xlsx')
    assert '[paragraf 2] Karar: bütçe onaylandı.' in s1.cikti
    assert '[Giderler!2. satır] 2026-10-01 | 120.5' in s2.cikti


def test_bozuk_ve_taranmis_belge_cokmez(ev):
    (ev / 'Downloads' / 'bozuk.pdf').write_bytes(b'%PDF-1.4 cop veri')
    (ev / 'Downloads' / 'bozuk.xlsx').write_bytes(b'PK\x03\x04 kirik zip')
    _pdf(ev / 'Downloads' / 'taranmis.pdf', [''])
    for ad in ('bozuk.pdf', 'bozuk.xlsx'):
        s = KAYIT['belge_oku'].calistir(_b(), yol=f'İndirilenler/{ad}')
        assert not s.basarili and s.hata_turu in ('bozuk', 'ocr_gerekli'), (ad, s.cikti)
    s = KAYIT['belge_oku'].calistir(_b(), yol='İndirilenler/taranmis.pdf')
    assert not s.basarili and s.hata_turu == 'ocr_gerekli'


def test_csv_ayirici_ve_cikti_siniri(ev):
    satirlar = ['ad;tutar'] + [f'kalem{i};{i}' for i in range(5000)]
    (ev / 'Documents' / 'liste.csv').write_text('\n'.join(satirlar), encoding='utf-8')
    s = KAYIT['belge_oku'].calistir(_b(), yol='Belgeler/liste.csv')
    assert s.basarili and '[2. satır] kalem0 | 0' in s.cikti and 'kısaltıldı' in s.cikti and len(s.cikti) < 13000
