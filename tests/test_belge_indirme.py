"""Documents, downloads and opening files."""
import http.server
import threading
import time

import pytest

from jarvis_core.araclar import KAYIT, Baglam, Iptal
from jarvis_core.izinler import SISTEM

def _pdf_yap(metin: str) -> bytes:
    """Minimal valid one-page PDF with a correct xref table."""
    akis = f'BT /F1 18 Tf 20 100 Td ({metin}) Tj ET'.encode()
    nesneler = [b'<</Type/Catalog/Pages 2 0 R>>', b'<</Type/Pages/Kids[3 0 R]/Count 1>>',
                b'<</Type/Page/Parent 2 0 R/MediaBox[0 0 300 144]/Contents 4 0 R/Resources<</Font<</F1 5 0 R>>>>>>',
                b'<</Length %d>>stream\n' % len(akis) + akis + b'\nendstream',
                b'<</Type/Font/Subtype/Type1/BaseFont/Helvetica>>']
    cikti, konumlar = b'%PDF-1.4\n', []
    for no, govde in enumerate(nesneler, 1):
        konumlar.append(len(cikti))
        cikti += b'%d 0 obj' % no + govde + b'endobj\n'
    xref = len(cikti)
    cikti += b'xref\n0 %d\n0000000000 65535 f \n' % (len(nesneler) + 1)
    cikti += b''.join(b'%010d 00000 n \n' % k for k in konumlar)
    cikti += b'trailer<</Size %d/Root 1 0 R>>\nstartxref\n%d\n%%%%EOF\n' % (len(nesneler) + 1, xref)
    return cikti


PDF = _pdf_yap('Havuz bakim raporu Ekim')


def _baglam():
    return Baglam(threading.Event())


def test_pdf_docx_xlsx_csv_okuma(ev):
    import docx
    import openpyxl
    (ev / 'Documents' / 'r.pdf').write_bytes(PDF)
    d = docx.Document(); d.add_paragraph('Klor seviyesi 1.5 ppm'); d.save(str(ev / 'Documents' / 'r.docx'))
    k = openpyxl.Workbook(); k.active.append(['Tarih', 'pH']); k.active.append(['16.10', 7.4]); k.save(str(ev / 'Documents' / 'r.xlsx'))
    (ev / 'Documents' / 'r.csv').write_text('ad;tutar\nkimyasal;350\n', encoding='utf-8')
    oku = KAYIT['belge_oku']
    assert 'Havuz bakim raporu' in oku.calistir(_baglam(), yol='Belgeler/r.pdf').cikti
    assert 'Klor seviyesi' in oku.calistir(_baglam(), yol='Belgeler/r.docx').cikti
    assert '7.4' in oku.calistir(_baglam(), yol='Belgeler/r.xlsx').cikti
    assert 'kimyasal | 350' in oku.calistir(_baglam(), yol='Belgeler/r.csv').cikti
    assert KAYIT['belge_oku'].guvenilmez_cikti


def test_rapor_docx_yazilir_ve_dogrulanir(ev):
    s = KAYIT['rapor_yaz'].calistir(_baglam(), yol='Belgeler/Ozet', baslik='PDF Özetleri',
                                     icerik='# Ekim\n- Klor normal\n- pH 7.4\nGenel durum iyi.')
    assert s.basarili and s.dogrulandi and (ev / 'Documents' / 'Ozet.docx').exists()
    oku = KAYIT['belge_oku'].calistir(_baglam(), yol='Belgeler/Ozet.docx').cikti
    assert 'PDF Özetleri' in oku and 'Klor normal' in oku


class _Sunucu(http.server.BaseHTTPRequestHandler):
    veri = b'x' * 300_000
    yalan_boyut = False
    yavas = False

    def do_GET(self):
        self.send_response(200)
        self.send_header('Content-Type', 'application/pdf')
        self.send_header('Content-Disposition', 'attachment; filename="fatura ekim.pdf"')
        self.send_header('Content-Length', str(len(self.veri) + (1000 if self.yalan_boyut else 0)))
        self.end_headers()
        for i in range(0, len(self.veri), 30_000):
            try:
                self.wfile.write(self.veri[i:i + 30_000])
            except OSError:
                return
            if self.yavas:
                time.sleep(0.3)

    def log_message(self, *a):
        pass


@pytest.fixture
def sunucu():
    s = http.server.ThreadingHTTPServer(('127.0.0.1', 0), _Sunucu)
    threading.Thread(target=s.serve_forever, daemon=True).start()
    yield f'http://127.0.0.1:{s.server_port}/dosya'
    _Sunucu.yalan_boyut = _Sunucu.yavas = False
    s.shutdown()


def test_indirme_dogrulanir_ve_uzerine_yazmaz(ev, sunucu):
    ilerleme = []
    b = Baglam(threading.Event(), bildir=ilerleme.append)
    s1 = KAYIT['dosya_indir'].calistir(b, adres=sunucu)
    assert s1.basarili and s1.dogrulandi and 'boyut doğrulandı' in s1.cikti
    assert (ev / 'Downloads' / 'fatura ekim.pdf').stat().st_size == 300_000
    s2 = KAYIT['dosya_indir'].calistir(b, adres=sunucu)
    assert (ev / 'Downloads' / 'fatura ekim (1).pdf').exists() and s2.dogrulandi
    assert not list((ev / 'Downloads').glob('*.indiriliyor'))


def test_eksik_inen_dosya_dogrulanmaz(ev, sunucu):
    _Sunucu.yalan_boyut = True
    s = KAYIT['dosya_indir'].calistir(_baglam(), adres=sunucu, ad='eksik.bin')
    assert not s.dogrulandi


def test_indirme_iptal_edilir_ve_yarim_dosya_kalmaz(ev, sunucu):
    _Sunucu.yavas = True
    iptal = threading.Event()
    sonuc = {}

    def calis():
        try:
            KAYIT['dosya_indir'].calistir(Baglam(iptal), adres=sunucu, ad='uzun.bin')
        except Iptal:
            sonuc['iptal'] = True
    t = threading.Thread(target=calis)
    t.start()
    time.sleep(0.8)
    iptal.set()
    t.join(5)
    assert sonuc.get('iptal')
    assert not list((ev / 'Downloads').glob('uzun*'))


def test_calistirilabilir_dosya_acmak_her_zaman_onay_ister(motor_yap, ev):
    iz = motor_yap([]).izinler
    ac = KAYIT['dosya_ac']
    iz.kural_ekle('dosya_ac', 'İndirilenler', 'izin')
    assert ac.risk_belirle({'yol': 'İndirilenler/kurulum.exe'}) == SISTEM
    assert iz.karar(ac, {'yol': 'İndirilenler/kurulum.exe'})[0] == 'sor'
    assert iz.karar(ac, {'yol': 'İndirilenler/fatura.pdf'})[0] == 'izin'


def test_indir_ve_ac_senaryosu_motorla(motor_yap, ev, sunucu, monkeypatch):
    """'Bu bağlantıdaki dosyayı indir, tamamlanınca aç.'"""
    import subprocess
    acilan = []
    monkeypatch.setattr(subprocess, 'Popen', lambda argv, **k: acilan.append(argv))
    m = motor_yap([[('dosya_indir', {'adres': sunucu})],
                   [('dosya_ac', {'yol': 'İndirilenler/fatura ekim.pdf'})], 'İndirip açtım efendim.'])
    s = m.calistir(f'{sunucu} bağlantısındaki dosyayı indir, tamamlanınca aç')
    adimlar = m.depo.adimlar(s.gorev_id)
    assert [a['arac'] for a in adimlar] == ['dosya_indir', 'dosya_ac']
    assert adimlar[0]['dogrulandi'] == 1 and acilan
    # opening could not be verified on this OS -> not reported as fully done
    assert s.durum == 'dogrulanamadi'


def test_pdfleri_ozetle_ve_rapor_olustur_senaryosu(motor_yap, ev):
    """'Bu klasördeki PDF'leri özetle ve rapor oluştur.'"""
    (ev / 'Documents' / 'Havuz').mkdir()
    (ev / 'Documents' / 'Havuz' / 'ekim.pdf').write_bytes(_pdf_yap('Ekim klor 1.5 ppm pH 7.4'))
    (ev / 'Documents' / 'Havuz' / 'kasim.pdf').write_bytes(_pdf_yap('Kasim klor 1.2 ppm pH 7.6'))
    m = motor_yap([
        [('dosya_ara', {'ad': '.pdf', 'klasor': 'Belgeler/Havuz'})],
        [('belge_oku', {'yol': 'Belgeler/Havuz/ekim.pdf'}), ('belge_oku', {'yol': 'Belgeler/Havuz/kasim.pdf'})],
        [('rapor_yaz', {'yol': 'Belgeler/Havuz/Özet Raporu.docx', 'baslik': 'Havuz Raporları Özeti',
                        'icerik': '- Ekim: klor 1.5 ppm, pH 7.4\n- Kasım: klor 1.2 ppm, pH 7.6'})],
        'İki PDF özetlendi ve rapor Belgeler/Havuz içine kaydedildi efendim.'])
    s = m.calistir("Belgeler/Havuz klasöründeki PDF'leri özetle ve rapor oluştur")
    assert s.durum == 'tamamlandi'
    assert (ev / 'Documents' / 'Havuz' / 'Özet Raporu.docx').exists()
    gorulen = '\n'.join(x['content'] for x in m.llm.gorulen[2] if x['role'] == 'tool')
    assert 'Ekim klor 1.5' in gorulen and 'Kasim klor 1.2' in gorulen and '<harici_veri' in gorulen
