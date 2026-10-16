"""Document reading (Stage F): PDF, DOCX, XLSX, CSV, MD/TXT.

Every piece of text keeps its location (page / paragraph / sheet!cell / row)
and the source file's sha256, so an answer or a generated table can point
back to where a value came from. Limits: file size, pages, rows and output
characters, so a huge or broken document cannot exhaust memory. Scanned PDFs
without a text layer are reported as such (OCR is not implemented yet)."""
import csv
import hashlib
from pathlib import Path

from .. import yollar
from ..izinler import OKUMA
from .temel import Sonuc, arac, metin

AZAMI_BOYUT = 50 * 1024 * 1024
AZAMI_KARAKTER = 12000


def _ozet(yol: Path) -> str:
    h = hashlib.sha256()
    with open(yol, 'rb') as f:
        for parca in iter(lambda: f.read(1 << 20), b''):
            h.update(parca)
    return h.hexdigest()


def pdf_parcalari(yol: Path, baglam=None, azami_sayfa=200):
    from pypdf import PdfReader
    okuyucu = PdfReader(str(yol))
    if okuyucu.is_encrypted:
        try:
            okuyucu.decrypt('')
        except Exception:
            raise ValueError('PDF parola korumalı.')
    for i, sayfa in enumerate(okuyucu.pages[:azami_sayfa], 1):
        if baglam:
            baglam.kontrol()
        yield f'sayfa {i}', (sayfa.extract_text() or '').strip()


def docx_parcalari(yol: Path, baglam=None):
    import docx
    belge = docx.Document(str(yol))
    for i, p in enumerate(belge.paragraphs, 1):
        if p.text.strip():
            yield f'paragraf {i}', p.text.strip()
    for t, tablo in enumerate(belge.tables, 1):
        for s, satir in enumerate(tablo.rows, 1):
            if baglam:
                baglam.kontrol()
            yield f'tablo {t} satır {s}', ' | '.join(c.text.strip() for c in satir.cells)


def xlsx_parcalari(yol: Path, baglam=None, azami_satir=2000):
    import openpyxl
    kitap = openpyxl.load_workbook(str(yol), read_only=True, data_only=True)
    try:
        for sayfa in kitap.worksheets:
            for s, satir in enumerate(sayfa.iter_rows(values_only=True), 1):
                if s > azami_satir:
                    yield f'{sayfa.title}', f'[… {azami_satir} satırdan sonrası okunmadı]'
                    break
                if baglam and s % 100 == 0:
                    baglam.kontrol()
                degerler = ['' if v is None else str(v) for v in satir]
                if any(degerler):
                    yield f'{sayfa.title}!{s}. satır', ' | '.join(degerler)
    finally:
        kitap.close()


def csv_parcalari(yol: Path, baglam=None, azami_satir=2000):
    with open(yol, encoding='utf-8-sig', errors='replace', newline='') as f:
        ornek = f.read(4096)
        f.seek(0)
        try:
            lehce = csv.Sniffer().sniff(ornek, delimiters=',;\t')
        except csv.Error:
            lehce = csv.excel
        for s, satir in enumerate(csv.reader(f, lehce), 1):
            if s > azami_satir:
                yield 'csv', f'[… {azami_satir} satırdan sonrası okunmadı]'
                break
            yield f'{s}. satır', ' | '.join(satir)


def metin_parcalari(yol: Path, baglam=None):
    for i, satir in enumerate(yol.read_text(encoding='utf-8', errors='replace').splitlines(), 1):
        if satir.strip():
            yield f'satır {i}', satir


OKUYUCULAR = {'.pdf': pdf_parcalari, '.docx': docx_parcalari, '.xlsx': xlsx_parcalari, '.xlsm': xlsx_parcalari,
              '.csv': csv_parcalari, '.txt': metin_parcalari, '.md': metin_parcalari}


def belge_oku_ham(yol: Path, baglam=None) -> dict:
    if yol.stat().st_size > AZAMI_BOYUT:
        raise ValueError(f'Dosya çok büyük ({yol.stat().st_size // 2**20} MB > {AZAMI_BOYUT // 2**20} MB).')
    okuyucu = OKUYUCULAR.get(yol.suffix.lower())
    if not okuyucu:
        raise ValueError(f'{yol.suffix} biçimi desteklenmiyor (desteklenen: {", ".join(OKUYUCULAR)}).')
    return {'yol': str(yol), 'sha256': _ozet(yol), 'parcalar': list(okuyucu(yol, baglam))}


@arac('belge_oku', 'PDF, Word (docx), Excel (xlsx), CSV veya metin belgesinin içeriğini konum bilgisiyle '
      '(sayfa, paragraf, hücre satırı) okur.',
      {'yol': metin('Belgenin yolu, ör. "İndirilenler/fatura.pdf"')}, ['yol'], risk=OKUMA, yan_etki=False,
      zaman_asimi=120, guvenilmez_cikti=True)
def belge_oku(baglam, yol):
    p = yollar.coz(yol)
    if not p.is_file():
        return Sonuc(False, f'Dosya bulunamadı: {p}', hata_turu='bulunamadi')
    try:
        sonuc = belge_oku_ham(p, baglam)
    except ValueError as e:
        return Sonuc(False, str(e), hata_turu='bicim')
    except Exception as e:
        return Sonuc(False, f'Belge okunamadı (bozuk olabilir): {e.__class__.__name__}: {str(e)[:160]}',
                     hata_turu='bozuk')
    dolu = [(k, t) for k, t in sonuc['parcalar'] if t]
    if p.suffix.lower() == '.pdf' and not dolu:
        return Sonuc(False, 'PDF\'de metin katmanı yok (taranmış belge). OCR henüz desteklenmiyor.',
                     hata_turu='ocr_gerekli', veri={'sha256': sonuc['sha256']})
    satirlar, toplam = [], 0
    for konum, icerik in dolu:
        satir = f'[{konum}] {icerik}'
        if toplam + len(satir) > AZAMI_KARAKTER:
            satirlar.append(f'[… kalan {len(dolu) - len(satirlar)} parça kısaltıldı]')
            break
        satirlar.append(satir)
        toplam += len(satir)
    return Sonuc(True, f'{p.name} (sha256 {sonuc["sha256"][:12]}):\n' + '\n'.join(satirlar),
                 veri={'yol': str(p), 'sha256': sonuc['sha256'], 'parca_sayisi': len(dolu)})
